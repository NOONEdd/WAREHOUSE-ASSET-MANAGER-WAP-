"""Writes the same small synthetic reference table with two independent .xlsx
producers (openpyxl and XlsxWriter), so the evaluation can check that ExcelJS
reads files it did not write itself. Invented data only. Usage:
    python3 -I make_reference.py OUTPUT_DIR
"""
import datetime
import os
import sys

import openpyxl
import xlsxwriter

HEADER = ["کد دارایی", "نام", "سریال", "تعداد", "تاریخ خرید", "فعال", "یادداشت"]
ROWS = [
    ["00123", "لپ‌تاپ نمونه", "SN-0001", 3, datetime.date(2024, 3, 15), True, "=cmd|' /C calc'!A0"],
    ["000045", "Монитор", None, 1, datetime.date(2023, 12, 1), False, "line1\nline2"],
    ["A-7", "پرینتر ۱۲۳", "۰۰۹۸۷", 2.5, datetime.date(2025, 1, 31), True, "مقدار"],
]


def with_openpyxl(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "دارایی‌ها"
    ws.append(HEADER)
    for row in ROWS:
        ws.append(row)
    for r in range(2, 5):
        ws.cell(row=r, column=1).number_format = "@"
        # openpyxl treats any text starting with "=" as a formula; keep it text.
        if isinstance(ws.cell(row=r, column=7).value, str):
            ws.cell(row=r, column=7).data_type = "s"
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
    ws["D5"] = "=SUM(D2:D4)"
    ws.merge_cells("A6:C6")
    ws["A6"] = "merged note"
    wb.save(path)


def with_xlsxwriter(path):
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet("دارایی‌ها")
    date = wb.add_format({"num_format": "yyyy-mm-dd"})
    ws.write_row(0, 0, HEADER)
    for r, row in enumerate(ROWS, start=1):
        for c, value in enumerate(row):
            if value is None:
                continue
            if isinstance(value, datetime.date):
                ws.write_datetime(r, c, datetime.datetime.combine(value, datetime.time()), date)
            elif isinstance(value, str):
                ws.write_string(r, c, value)
            else:
                ws.write(r, c, value)
    ws.write_formula(4, 3, "=SUM(D2:D4)", None, 6.5)
    ws.merge_range("A6:C6", "merged note")
    wb.close()


if __name__ == "__main__":
    out = sys.argv[1]
    os.makedirs(out, exist_ok=True)
    with_openpyxl(os.path.join(out, "reference-openpyxl.xlsx"))
    with_xlsxwriter(os.path.join(out, "reference-xlsxwriter.xlsx"))
    print("written")
